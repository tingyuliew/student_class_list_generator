import streamlit as st
from docx import Document
from copy import deepcopy
import random
import math
import tempfile
import os
from docx.shared import Inches
from PIL import Image
import tempfile
import zipfile
import os
from lxml import etree
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

st.set_page_config(
page_title="Student Group Generator",
page_icon="📋",
layout="wide"
)

st.title("📋 Student Group Generator")
st.write(
"Hi Dad! This website will randomise students into groups and generate a new Word document.\n\n" 
"Please only upload 1 Word document at a time.\n"
"The document should contain a table with student information, including their name, student number, and image."
)

uploaded_doc = st.file_uploader(
"Upload Word Document",
type=["docx"]
)

group_size = st.number_input(
"Number of students per group",
min_value=1,
value=4,
step=1
)

if st.button("Generate Groups"):


    if uploaded_doc is None:
        st.error("Please upload a Word document.")
        st.stop()

    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".docx"
    ) as tmp:

        tmp.write(uploaded_doc.read())
        input_path = tmp.name

    try:

        doc = Document(input_path)

        if len(doc.tables) == 0:
            st.error("No tables found in document.")
            st.stop()

        source_table = doc.tables[0]

        ## Extract class header
        header_text = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                header_text.append(text)
        header_text = "\n".join(header_text)

        ### Extract student information from the table

        students = []
        
        for row in source_table.rows:
            for cell in row.cells:
                text = cell.text.strip()

                if not text:
                    continue

                lines = [x.strip() for x in text.split("\n") if x.strip()]

                if len(lines) < 2:
                    continue

                student_number = lines[-1]
                student_name = " ".join(lines[:-1])

                image_path = None

                # Search for image relationship
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:

                        xml = run._element.xml

                        if "r:embed" not in xml:
                            continue

                        tree = etree.fromstring(run._element.xml.encode())

                        embeds = tree.xpath(
                            './/*[local-name()="blip"]'
                        )

                        if len(embeds) == 0:
                            continue

                        rId = embeds[0].get(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                        )

                        image_part = doc.part.related_parts[rId]

                        image_bytes = image_part.blob

                        fd, image_path = tempfile.mkstemp(
                            suffix=".jpg"
                        )

                        with os.fdopen(fd, "wb") as f:
                            f.write(image_bytes)

                        break

                students.append({
                    "name": student_name,
                    "student_number": student_number,
                    "image_path": image_path
                })

        total_students = len(students)

        ### now group the students randomly according to the specified group size

        if total_students == 0:
            st.error("No student entries detected.")
            st.stop()

        random.shuffle(students)

        num_groups = math.ceil(
            total_students / group_size
        )

        groups = []

        for i in range(num_groups):

            start = i * group_size
            end = start + group_size

            groups.append(
                students[start:end]
            )

        ### Generate output document

        output_doc = Document()

        # make narrow margins 
        section = output_doc.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

        # add initial class text
        head = output_doc.add_paragraph()
        run = head.add_run(header_text)
        run.font.name = "Verdana (Body)"
        run.font.size = Pt(12)
        run.bold = True


        original_cols = len(source_table.columns)

        for idx, group in enumerate(groups, start=1):

            output_doc.add_heading(
                f"Group {idx}",
                level=2
            )

            rows_needed = math.ceil(
                len(group) / original_cols
            )

            table = output_doc.add_table(
                rows=rows_needed,
                cols=original_cols
            )
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            student_idx = 0

            for r in range(rows_needed):

                for c in range(original_cols):

                    if student_idx >= len(group):
                        continue

                    student = group[student_idx]

                    cell = table.cell(r, c) 
                    cell.text = ""

                    # Add image 
                    if student["image_path"]: 
                        p = cell.paragraphs[0] 
                        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                        run = p.add_run() 
                        run.add_picture(
                            student["image_path"], 
                            width=Inches(0.9))
                    
                    # Add name 
                    name_para = cell.add_paragraph()
                    name_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                    run = name_para.add_run(
                        student["name"])
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)

                    # Add student number 
                    num_para = cell.add_paragraph()
                    num_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                    run = num_para.add_run(
                        student["student_number"])
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)

                    student_idx += 1

        original_name = os.path.splitext(uploaded_doc.name)[0]

        output_path = f"{original_name}_grouped.docx"

        output_doc.save(output_path)

        with open(output_path, "rb") as file:

            st.success(
                f"Generated {num_groups} groups."
            )

            st.download_button(
                label="Download Grouped Document",
                data=file,
                file_name=output_path,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

    except Exception as e:
        st.error(str(e))

    finally:
        os.remove(input_path)
